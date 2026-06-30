#!/usr/bin/env python3
"""
Job Search Agent — Phase 1
Scans USA + Dubai job portals for Data Engineer / SQL Developer roles.
Flags visa sponsorship. Extracts contact emails where available.
Output: reports/jobs_YYYY-MM-DD.xlsx

Usage:
    python run_search.py               # search last 7 days
    python run_search.py --days 3      # search last 3 days
    python run_search.py --no-dubai    # USA only
    python run_search.py --no-usa      # Dubai only
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta
from pathlib import Path
from typing import List, Optional
from urllib.parse import quote_plus, urljoin

import requests
from bs4 import BeautifulSoup

# ── Logging ────────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-7s | %(message)s",
    datefmt="%H:%M:%S",
    handlers=[logging.StreamHandler(sys.stdout)],
)
log = logging.getLogger(__name__)

# ── Config ─────────────────────────────────────────────────────────────────────
USA_QUERIES = [
    "Data Engineer AWS",
    "Data Engineer GCP",
    "SQL Developer",
    "Senior Data Engineer",
    "Data Engineer Databricks",
    "Data Engineer Snowflake",
    "PL SQL Developer",
    "Data Engineer",
]

DUBAI_QUERIES = [
    "Data Engineer",
    "SQL Developer",
    "Data Engineer AWS GCP",
    "Senior Data Engineer",
]

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

# ── Patterns ───────────────────────────────────────────────────────────────────
_SPONSOR_YES = re.compile(
    r"(h[-\s]?1[-\s]?b\s*(sponsor|visa|transfer|cap)?|"
    r"visa\s+sponsor(ship)?|will\s+sponsor|sponsorship\s+(available|provided|offered)|"
    r"open\s+to\s+sponsor|able\s+to\s+sponsor|"
    r"sponsor\s+work\s+(visa|permit)|"
    r"work\s+permit\s+sponsor|relocation\s+(package|support|assistance)|"
    r"we\s+sponsor|company\s+sponsored\s+visa|"
    r"employment\s+(authorization|visa)\s+sponsor)",
    re.IGNORECASE,
)
_SPONSOR_NO = re.compile(
    r"(no\s+(visa\s+)?sponsor(ship)?|"
    r"must\s+be\s+(authorized|eligible|a\s+(us\s+)?citizen)|"
    r"authorized\s+to\s+work\s+in\s+the\s+u\.?s\.?\s+without|"
    r"no\s+h[-\s]?1[-\s]?b|"
    r"u\.?s\.?\s+citizen(s)?\s+(or|/)\s+permanent\s+resident|"
    r"(green\s+card|gc)\s+holder\s+only|"
    r"not\s+(able|willing)\s+to\s+sponsor)",
    re.IGNORECASE,
)
_EMAIL = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")


# ── Data model ─────────────────────────────────────────────────────────────────
class Job:
    __slots__ = [
        "title", "company", "location", "region", "source",
        "url", "email", "sponsorship", "posted_date", "description",
    ]

    def __init__(
        self,
        title: str,
        company: str,
        location: str,
        region: str,       # "USA" | "Dubai"
        source: str,
        url: str,
        description: str = "",
        posted_date: Optional[datetime] = None,
    ):
        self.title = (title or "Unknown").strip()
        self.company = (company or "Unknown").strip()
        self.location = (location or "").strip()
        self.region = region
        self.source = source
        self.url = (url or "").strip()
        self.description = description or ""
        self.posted_date = posted_date

        # Extract first email found in description
        emails = _EMAIL.findall(self.description)
        # Filter out image/asset file extensions mistaken for emails
        real_emails = [e for e in emails if not re.search(r"\.(png|jpg|gif|svg|css|js)$", e, re.I)]
        self.email = real_emails[0] if real_emails else ""

        # Determine sponsorship
        if _SPONSOR_YES.search(self.description):
            self.sponsorship = "Yes"
        elif _SPONSOR_NO.search(self.description):
            self.sponsorship = "No"
        else:
            self.sponsorship = "Not mentioned"

    def to_row(self) -> list:
        date_str = self.posted_date.strftime("%Y-%m-%d") if self.posted_date else ""
        return [
            self.title, self.company, self.location, self.region,
            self.source, self.url, self.email, self.sponsorship, date_str,
        ]


# ── Helpers ────────────────────────────────────────────────────────────────────
def _parse_relative_date(raw: str) -> Optional[datetime]:
    """Parse 'X hours ago', 'X days ago', ISO strings, etc."""
    if not raw:
        return None
    raw = raw.strip().lower()
    if "just now" in raw or "today" in raw:
        return datetime.utcnow()
    m = re.search(r"(\d+)\s+hour", raw)
    if m:
        return datetime.utcnow() - timedelta(hours=int(m.group(1)))
    m = re.search(r"(\d+)\s+day", raw)
    if m:
        return datetime.utcnow() - timedelta(days=int(m.group(1)))
    m = re.search(r"(\d+)\s+week", raw)
    if m:
        return datetime.utcnow() - timedelta(weeks=int(m.group(1)))
    m = re.search(r"(\d+)\s+month", raw)
    if m:
        return datetime.utcnow() - timedelta(days=int(m.group(1)) * 30)
    try:
        return datetime.fromisoformat(raw[:19].replace("z", ""))
    except ValueError:
        pass
    return None


def _get(url: str, timeout: int = 15) -> Optional[requests.Response]:
    """HTTP GET with polite retry."""
    for attempt in range(3):
        try:
            r = requests.get(url, headers=HEADERS, timeout=timeout)
            if r.status_code == 200:
                return r
            if r.status_code in (429, 503):
                wait = 10 * (attempt + 1)
                log.warning(f"Rate limit {r.status_code} — waiting {wait}s")
                time.sleep(wait)
            else:
                log.debug(f"HTTP {r.status_code} for {url}")
                return None
        except requests.RequestException as e:
            log.debug(f"Request error ({url}): {e}")
            time.sleep(5)
    return None


# ══════════════════════════════════════════════════════════════════════════════
# USA SCRAPERS
# ══════════════════════════════════════════════════════════════════════════════

def _jobspy_call(site: str, query: str, location: str, days: int) -> list:
    """Call jobspy with progressive fallback for different installed versions."""
    try:
        from jobspy import scrape_jobs
    except ImportError:
        return []

    hours_old = days * 24
    base = dict(
        site_name=[site],
        search_term=query,
        location=location,
        results_wanted=50,
        country_indeed="USA" if "United States" in location else "AE",
    )
    for extras in [
        {"hours_old": hours_old, "verbose": 0},
        {"hours_old": hours_old},
        {},
    ]:
        try:
            df = scrape_jobs(**base, **extras)
            return [] if (df is None or df.empty) else df.to_dict("records")
        except TypeError:
            continue
        except Exception as e:
            log.warning(f"[jobspy:{site}] {query!r}: {e}")
            return []
    return []


def scrape_jobspy_usa(query: str, days: int) -> List[Job]:
    """LinkedIn, Indeed, ZipRecruiter, Glassdoor via jobspy — USA."""
    jobs: List[Job] = []
    site_labels = {
        "linkedin": "LinkedIn",
        "indeed": "Indeed",
        "zip_recruiter": "ZipRecruiter",
        "glassdoor": "Glassdoor",
    }
    for site, label in site_labels.items():
        rows = _jobspy_call(site, query, "United States", days)
        for row in rows:
            url = str(row.get("job_url") or row.get("job_url_direct") or "")
            if not url:
                continue
            raw_date = row.get("date_posted")
            posted = None
            if raw_date:
                try:
                    if isinstance(raw_date, datetime):
                        posted = raw_date
                    else:
                        posted = _parse_relative_date(str(raw_date))
                except Exception:
                    pass
            jobs.append(Job(
                title=str(row.get("title") or ""),
                company=str(row.get("company") or ""),
                location=str(row.get("location") or "United States"),
                region="USA",
                source=label,
                url=url,
                description=str(row.get("description") or ""),
                posted_date=posted,
            ))
    log.info(f"[jobspy-USA] {query!r} → {len(jobs)} jobs")
    return jobs


def scrape_dice(query: str) -> List[Job]:
    """Dice.com — tech-focused US job board."""
    jobs: List[Job] = []
    url = (
        f"https://www.dice.com/jobs?q={quote_plus(query)}"
        f"&countryCode=US&radius=30&radiusUnit=mi&page=1&pageSize=50"
        f"&filters.postedDate=SEVEN"
    )
    resp = _get(url)
    if not resp:
        return []

    import json
    soup = BeautifulSoup(resp.text, "html.parser")

    # Try __NEXT_DATA__ embedded JSON
    script = soup.find("script", id="__NEXT_DATA__")
    if script and script.string:
        try:
            nd = json.loads(script.string)
            props = nd.get("props", {}).get("pageProps", {})

            # Try multiple known paths
            raw_jobs = (
                props.get("initialState", {}).get("jobs", {}).get("jobs", [])
                or props.get("jobs", [])
                or props.get("searchResults", {}).get("jobs", [])
                or []
            )

            # React Query / dehydratedState path
            if not raw_jobs:
                for entry in props.get("dehydratedState", {}).get("queries", []):
                    data = entry.get("state", {}).get("data", {})
                    if isinstance(data, dict):
                        for v in data.values():
                            if isinstance(v, list) and v and isinstance(v[0], dict):
                                if "title" in v[0] or "jobTitle" in v[0]:
                                    raw_jobs.extend(v)
                    elif isinstance(data, list) and data and isinstance(data[0], dict):
                        raw_jobs.extend(data)

            for job in raw_jobs:
                title   = str(job.get("title") or job.get("jobTitle") or "Unknown")
                company = job.get("advertiser") or job.get("company") or {}
                company = company.get("name", "Unknown") if isinstance(company, dict) else str(company)
                loc     = job.get("location") or {}
                loc     = loc.get("displayName", "United States") if isinstance(loc, dict) else str(loc)
                jid     = job.get("id") or job.get("jobId") or ""
                jurl    = job.get("applyUrl") or (f"https://www.dice.com/job-detail/{jid}" if jid else "")
                snippet = str(job.get("summary") or job.get("jobDescription") or "")
                date_raw = str(job.get("postedDate") or "")
                jobs.append(Job(
                    title=title, company=company, location=loc,
                    region="USA", source="Dice", url=jurl,
                    description=snippet,
                    posted_date=_parse_relative_date(date_raw),
                ))
        except Exception as e:
            log.debug(f"[Dice] JSON parse error: {e}")

    # Fallback: HTML card scraping
    if not jobs:
        for card in soup.select("dhi-search-card, [data-cy='card-title-link']"):
            try:
                title_el = (
                    card.find(attrs={"data-cy": "card-title-link"})
                    or card.find(["h2", "h3"])
                )
                if not title_el:
                    continue
                title = title_el.get_text(strip=True)
                link_el = card.find("a", href=True)
                href = link_el["href"] if link_el else ""
                jurl = f"https://www.dice.com{href}" if href.startswith("/") else href
                comp_el = card.find(attrs={"data-cy": "search-result-company-name"})
                company = comp_el.get_text(strip=True) if comp_el else "Unknown"
                loc_el  = card.find(attrs={"data-cy": "search-result-location"})
                location = loc_el.get_text(strip=True) if loc_el else "United States"
                date_el = card.find(attrs={"data-cy": "card-posted-date"})
                date_raw = date_el.get_text(strip=True) if date_el else ""
                jobs.append(Job(
                    title=title, company=company, location=location,
                    region="USA", source="Dice", url=jurl,
                    posted_date=_parse_relative_date(date_raw),
                ))
            except Exception:
                pass

    log.info(f"[Dice] {query!r} → {len(jobs)} jobs")
    return jobs


# ══════════════════════════════════════════════════════════════════════════════
# DUBAI SCRAPERS
# ══════════════════════════════════════════════════════════════════════════════

def scrape_jobspy_dubai(query: str, days: int) -> List[Job]:
    """LinkedIn UAE via jobspy."""
    jobs: List[Job] = []
    rows = _jobspy_call("linkedin", query, "Dubai, United Arab Emirates", days)
    for row in rows:
        url = str(row.get("job_url") or row.get("job_url_direct") or "")
        if not url:
            continue
        raw_date = row.get("date_posted")
        posted = None
        if raw_date:
            try:
                posted = raw_date if isinstance(raw_date, datetime) else _parse_relative_date(str(raw_date))
            except Exception:
                pass
        loc = str(row.get("location") or "Dubai, UAE")
        if not any(k in loc.lower() for k in ("uae", "dubai", "abu dhabi", "sharjah", "emirates")):
            continue  # skip non-UAE results
        jobs.append(Job(
            title=str(row.get("title") or ""),
            company=str(row.get("company") or ""),
            location=loc,
            region="Dubai",
            source="LinkedIn UAE",
            url=url,
            description=str(row.get("description") or ""),
            posted_date=posted,
        ))
    log.info(f"[LinkedIn-UAE] {query!r} → {len(jobs)} jobs")
    return jobs


def scrape_bayt(query: str) -> List[Job]:
    """Bayt.com — leading UAE/Middle East job portal."""
    jobs: List[Job] = []
    slug = query.lower().replace(" ", "-")
    urls_to_try = [
        f"https://www.bayt.com/en/uae/jobs/{slug}-jobs/",
        f"https://www.bayt.com/en/uae/jobs/?q={quote_plus(query)}&l=1",
    ]

    resp = None
    for url in urls_to_try:
        resp = _get(url)
        if resp:
            break
    if not resp:
        log.warning(f"[Bayt] Could not reach site for {query!r}")
        return []

    soup = BeautifulSoup(resp.text, "html.parser")

    # Bayt uses <li data-js-job> for job cards
    cards = soup.select("li[data-js-job]") or soup.select("article.jb-item")

    for card in cards:
        try:
            title_el = (
                card.find("h2", class_=re.compile(r"jb-title", re.I))
                or card.find("a", class_=re.compile(r"jb-title", re.I))
                or card.find(["h2", "h3"])
            )
            if not title_el:
                continue
            title = title_el.get_text(strip=True)

            link_el = title_el if title_el.name == "a" else title_el.find("a")
            if not link_el:
                link_el = card.find("a", href=re.compile(r"/job/"))
            href = link_el.get("href", "") if link_el else ""
            jurl = f"https://www.bayt.com{href}" if href.startswith("/") else href

            comp_el = (
                card.find(class_=re.compile(r"jb-company|company", re.I))
                or card.find("b")
            )
            company = comp_el.get_text(strip=True) if comp_el else "Unknown"

            loc_el = card.find(class_=re.compile(r"location|city|country", re.I))
            location = loc_el.get_text(strip=True) if loc_el else "Dubai, UAE"

            date_el = card.find(class_=re.compile(r"date|posted|age", re.I))
            date_raw = date_el.get_text(strip=True) if date_el else ""

            snippet_el = card.find(class_=re.compile(r"desc|summary|snippet", re.I))
            snippet = snippet_el.get_text(strip=True) if snippet_el else ""

            jobs.append(Job(
                title=title, company=company, location=location,
                region="Dubai", source="Bayt.com", url=jurl,
                description=snippet,
                posted_date=_parse_relative_date(date_raw),
            ))
        except Exception as e:
            log.debug(f"[Bayt] card parse error: {e}")

    log.info(f"[Bayt] {query!r} → {len(jobs)} jobs")
    return jobs


def scrape_gulftablent(query: str) -> List[Job]:
    """GulfTalent — premium UAE/Gulf job board."""
    jobs: List[Job] = []
    url = f"https://www.gulftalent.com/jobs?query={quote_plus(query)}&country=ae"
    resp = _get(url)
    if not resp:
        log.warning(f"[GulfTalent] Could not reach site for {query!r}")
        return []

    soup = BeautifulSoup(resp.text, "html.parser")
    cards = soup.select(".job-listing-item, .job-item, article.job, li.job")

    for card in cards:
        try:
            title_el = card.find(["h2", "h3"]) or card.find("a", class_=re.compile(r"title", re.I))
            if not title_el:
                continue
            title = title_el.get_text(strip=True)
            link_el = card.find("a", href=True)
            href = link_el["href"] if link_el else ""
            jurl = urljoin("https://www.gulftalent.com", href)
            comp_el = card.find(class_=re.compile(r"company|employer", re.I))
            company = comp_el.get_text(strip=True) if comp_el else "Unknown"
            loc_el = card.find(class_=re.compile(r"location|city", re.I))
            location = loc_el.get_text(strip=True) if loc_el else "UAE"
            date_el = card.find(class_=re.compile(r"date|posted|ago", re.I))
            date_raw = date_el.get_text(strip=True) if date_el else ""
            jobs.append(Job(
                title=title, company=company, location=location,
                region="Dubai", source="GulfTalent", url=jurl,
                posted_date=_parse_relative_date(date_raw),
            ))
        except Exception:
            pass

    log.info(f"[GulfTalent] {query!r} → {len(jobs)} jobs")
    return jobs


def scrape_naukrigulf(query: str) -> List[Job]:
    """Naukrigulf — large UAE job board."""
    jobs: List[Job] = []
    url = f"https://www.naukrigulf.com/{quote_plus(query.lower().replace(' ', '-'))}-jobs"
    resp = _get(url)
    if not resp:
        log.warning(f"[Naukrigulf] Could not reach site for {query!r}")
        return []

    soup = BeautifulSoup(resp.text, "html.parser")
    cards = soup.select(".ni-job-tuple, .job-listing, article.job-post")

    for card in cards:
        try:
            title_el = card.find(["h2", "h3", "a"], class_=re.compile(r"title|job-title", re.I))
            if not title_el:
                continue
            title = title_el.get_text(strip=True)
            link_el = title_el if title_el.name == "a" else title_el.find("a")
            href = link_el.get("href", "") if link_el else ""
            jurl = urljoin("https://www.naukrigulf.com", href)
            comp_el = card.find(class_=re.compile(r"company|employer|org", re.I))
            company = comp_el.get_text(strip=True) if comp_el else "Unknown"
            loc_el = card.find(class_=re.compile(r"location|city|loc", re.I))
            location = loc_el.get_text(strip=True) if loc_el else "UAE"
            date_el = card.find(class_=re.compile(r"date|posted|age", re.I))
            date_raw = date_el.get_text(strip=True) if date_el else ""
            jobs.append(Job(
                title=title, company=company, location=location,
                region="Dubai", source="Naukrigulf", url=jurl,
                posted_date=_parse_relative_date(date_raw),
            ))
        except Exception:
            pass

    log.info(f"[Naukrigulf] {query!r} → {len(jobs)} jobs")
    return jobs


# ══════════════════════════════════════════════════════════════════════════════
# DEDUPLICATION
# ══════════════════════════════════════════════════════════════════════════════

def deduplicate(jobs: List[Job]) -> List[Job]:
    """Remove duplicates by URL first, then by title+company."""
    seen_urls: set = set()
    seen_tc: set = set()
    unique: List[Job] = []
    for j in jobs:
        url_key = j.url.rstrip("/").lower()
        tc_key = (j.title.lower().strip(), j.company.lower().strip())
        if url_key and url_key in seen_urls:
            continue
        if tc_key in seen_tc:
            continue
        seen_urls.add(url_key)
        seen_tc.add(tc_key)
        unique.append(j)
    return unique


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL EXPORT
# ══════════════════════════════════════════════════════════════════════════════

def save_excel(jobs: List[Job], output_path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.error("openpyxl not installed. Run: pip install openpyxl")
        sys.exit(1)

    wb = Workbook()

    # ── Jobs sheet ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Jobs"

    COLS = [
        ("Job Title",      40),
        ("Company",        25),
        ("Location",       22),
        ("Region",         10),
        ("Source",         16),
        ("Job Link",       14),
        ("Contact Email",  30),
        ("Sponsorship",    16),
        ("Posted Date",    13),
    ]

    # Header row
    HDR_FILL = PatternFill("solid", fgColor="1F4E79")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
    for col_idx, (name, width) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # Sponsorship colors
    SPONSOR_FILL = {
        "Yes":           PatternFill("solid", fgColor="C6EFCE"),
        "No":            PatternFill("solid", fgColor="FFC7CE"),
        "Not mentioned": PatternFill("solid", fgColor="FFEB9C"),
    }
    SPONSOR_FONT = {
        "Yes":           Font(color="276221", bold=True),
        "No":            Font(color="9C0006", bold=True),
        "Not mentioned": Font(color="7D6608"),
    }
    REGION_FILL = {
        "USA":   PatternFill("solid", fgColor="DDEEFF"),
        "Dubai": PatternFill("solid", fgColor="FFF0DD"),
    }

    thin = Side(border_style="thin", color="D0D0D0")
    row_border = Border(bottom=thin)

    url_font      = Font(color="0563C1", underline="single")
    email_font    = Font(color="0563C1")
    default_align = Alignment(vertical="top", wrap_text=False)

    for row_idx, job in enumerate(jobs, 2):
        row_data = job.to_row()
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = default_align
            cell.border = row_border

            # Region column (4) — color whole row lightly
            if col_idx == 4 and val in REGION_FILL:
                pass  # applied below

            # Job Link (6) — hyperlink
            if col_idx == 6 and val:
                cell.hyperlink = val
                cell.value = "View Job"
                cell.font = url_font

            # Contact Email (7)
            elif col_idx == 7 and val:
                cell.font = email_font

            # Sponsorship (8)
            elif col_idx == 8 and val in SPONSOR_FILL:
                cell.fill = SPONSOR_FILL[val]
                cell.font = SPONSOR_FONT[val]
                cell.alignment = Alignment(horizontal="center", vertical="top")

        # Light row tinting by region
        region = job.region
        if region in REGION_FILL:
            for col_idx in (1, 2, 3, 4, 5, 9):
                ws.cell(row=row_idx, column=col_idx).fill = REGION_FILL[region]

    # ── Summary sheet ───────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Job Search Summary"
    ws2["A1"].font = Font(bold=True, size=14, color="1F4E79")

    usa_jobs    = [j for j in jobs if j.region == "USA"]
    dubai_jobs  = [j for j in jobs if j.region == "Dubai"]
    sponsor_yes = [j for j in jobs if j.sponsorship == "Yes"]
    sponsor_no  = [j for j in jobs if j.sponsorship == "No"]

    summary_rows = [
        ("Run Date",                    datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")),
        ("Total Jobs Found",            len(jobs)),
        ("USA Jobs",                    len(usa_jobs)),
        ("Dubai Jobs",                  len(dubai_jobs)),
        ("Sponsorship: Yes",            len(sponsor_yes)),
        ("Sponsorship: No",             len(sponsor_no)),
        ("Sponsorship: Not mentioned",  len(jobs) - len(sponsor_yes) - len(sponsor_no)),
        ("Jobs with Contact Email",     sum(1 for j in jobs if j.email)),
    ]

    # Source breakdown
    from collections import Counter
    source_counts = Counter(j.source for j in jobs)
    summary_rows.append(("", ""))
    summary_rows.append(("Source", "Count"))
    for src, cnt in sorted(source_counts.items(), key=lambda x: -x[1]):
        summary_rows.append((src, cnt))

    LBL_FONT = Font(bold=True)
    for r_idx, (label, value) in enumerate(summary_rows, 3):
        ws2.cell(row=r_idx, column=1, value=label).font = LBL_FONT
        ws2.cell(row=r_idx, column=2, value=value)

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 20

    # ── Save ────────────────────────────────────────────────────
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════

def parse_args():
    p = argparse.ArgumentParser(description="Job Search Agent — Phase 1")
    p.add_argument("--days",     type=int, default=7,  help="Search window in days (default: 7)")
    p.add_argument("--no-usa",   action="store_true",   help="Skip USA sources")
    p.add_argument("--no-dubai", action="store_true",   help="Skip Dubai sources")
    return p.parse_args()


def main():
    args = parse_args()
    all_jobs: List[Job] = []

    log.info("=" * 60)
    log.info(f"Job Search Agent — Phase 1  |  {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    log.info(f"Window: last {args.days} days")
    log.info("=" * 60)

    # ── USA ──────────────────────────────────────────────────────
    if not args.no_usa:
        log.info("\n--- USA Sources ---")
        with ThreadPoolExecutor(max_workers=4) as pool:
            futures = []
            for q in USA_QUERIES:
                futures.append(pool.submit(scrape_jobspy_usa, q, args.days))
                futures.append(pool.submit(scrape_dice, q))
            for f in as_completed(futures):
                try:
                    all_jobs.extend(f.result())
                except Exception as e:
                    log.error(f"USA scraper error: {e}")

    # ── Dubai ────────────────────────────────────────────────────
    if not args.no_dubai:
        log.info("\n--- Dubai Sources ---")
        with ThreadPoolExecutor(max_workers=3) as pool:
            futures = []
            for q in DUBAI_QUERIES:
                futures.append(pool.submit(scrape_jobspy_dubai, q, args.days))
                futures.append(pool.submit(scrape_bayt, q))
                futures.append(pool.submit(scrape_gulftablent, q))
                futures.append(pool.submit(scrape_naukrigulf, q))
            for f in as_completed(futures):
                try:
                    all_jobs.extend(f.result())
                except Exception as e:
                    log.error(f"Dubai scraper error: {e}")

    # ── Dedup + Stats ────────────────────────────────────────────
    log.info(f"\nRaw total   : {len(all_jobs)} jobs")
    unique = deduplicate(all_jobs)
    log.info(f"After dedup : {len(unique)} unique jobs")

    usa_count    = sum(1 for j in unique if j.region == "USA")
    dubai_count  = sum(1 for j in unique if j.region == "Dubai")
    sponsor_yes  = sum(1 for j in unique if j.sponsorship == "Yes")
    with_email   = sum(1 for j in unique if j.email)
    log.info(f"USA: {usa_count}  |  Dubai: {dubai_count}  |  Sponsorship mentioned: {sponsor_yes}  |  With email: {with_email}")

    if not unique:
        log.warning("No jobs found. Check your internet connection and try: pip install --upgrade python-jobspy")
        return

    # ── Sort: Sponsorship Yes first, then by Region, then by Source ──
    unique.sort(key=lambda j: (
        0 if j.sponsorship == "Yes" else (1 if j.sponsorship == "Not mentioned" else 2),
        j.region,
        j.source,
        j.title,
    ))

    # ── Save Excel ───────────────────────────────────────────────
    date_str     = datetime.utcnow().strftime("%Y-%m-%d")
    output_path  = Path(__file__).parent / "reports" / f"jobs_{date_str}.xlsx"
    save_excel(unique, output_path)
    log.info(f"\nExcel saved → {output_path}")
    log.info(f"Open the file and check the 'Jobs' and 'Summary' sheets.")
    log.info("=" * 60)


if __name__ == "__main__":
    main()
